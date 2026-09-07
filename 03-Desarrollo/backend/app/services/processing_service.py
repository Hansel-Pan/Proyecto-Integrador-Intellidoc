import os
import uuid
import logging
from typing import List, Optional
from sqlalchemy.orm import Session

import pdfplumber
from docx import Document as DocxDocument
from sentence_transformers import SentenceTransformer

from app.core.config import settings
from app.models import Documento, Fragmento, LogProcesamiento
from app.services.vector_store import get_document_collection

logger = logging.getLogger(__name__)


class ProcessingService:
    def __init__(self, db: Session):
        self.db = db
        self.embedding_model = None
        self.vector_store = None

    def _get_embedding_model(self) -> SentenceTransformer:
        if self.embedding_model is None:
            self.embedding_model = SentenceTransformer(settings.EMBEDDING_MODEL)
        return self.embedding_model

    def _get_chroma_collection(self):
        return get_document_collection()

    def extract_text(self, file_path: str, formato: str) -> str:
        """Extraer texto según el formato del archivo."""
        try:
            if formato == "pdf":
                return self._extract_pdf(file_path)
            elif formato == "docx":
                return self._extract_docx(file_path)
            elif formato == "txt":
                return self._extract_txt(file_path)
            else:
                raise ValueError(f"Formato no soportado: {formato}")
        except Exception as e:
            logger.error(f"Error extrayendo texto de {file_path}: {e}")
            raise

    def _extract_pdf(self, file_path: str) -> str:
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return "\n\n".join(text_parts)

    def _extract_docx(self, file_path: str) -> str:
        doc = DocxDocument(file_path)
        text_parts = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    text_parts.append(" | ".join(cells))
        return "\n".join(text_parts)

    def _extract_txt(self, file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    def chunk_text(self, text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
        """Dividir texto en chunks con solapamiento."""
        words = text.split()
        chunks = []
        
        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append(chunk)
            if i + chunk_size >= len(words):
                break
        return chunks

    def generate_embeddings(self, texts: List[str]) -> List[List[float]]:
        model = self._get_embedding_model()
        embeddings = model.encode(texts, convert_to_tensor=False)
        return embeddings.tolist()

    def store_embeddings(self, documento_id: uuid.UUID, chunks: List[str], embeddings: List[List[float]]) -> List[Fragmento]:
        collection = self._get_chroma_collection()
        fragmentos = []
        
        for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
            vector_id = f"{documento_id}_{i}"
            
            collection.add(
                ids=[vector_id],
                embeddings=[embedding],
                documents=[chunk],
                metadatas=[{"documento_id": str(documento_id), "orden": i}]
            )
            
            fragmento = Fragmento(
                documento_id=documento_id,
                texto=chunk,
                vector_id=vector_id,
                orden=i
            )
            self.db.add(fragmento)
            fragmentos.append(fragmento)
        
        self.db.commit()
        for f in fragmentos:
            self.db.refresh(f)
        return fragmentos

    def process_document_sync(self, documento_id: uuid.UUID) -> bool:
        """Pipeline completo de procesamiento de un documento (SÍNCRONO)."""
        documento = self.db.query(Documento).filter(Documento.id == documento_id).first()
        if not documento:
            self.log_error(documento_id, "documento_no_encontrado", "Documento no encontrado en BD")
            return False

        try:
            self.update_status(documento_id, "procesando")
            self.log_info(documento_id, "inicio_procesamiento", "Iniciando pipeline de procesamiento")

            # 1. Extraer texto
            self.log_info(documento_id, "extraccion_texto", "Extrayendo texto del archivo")
            texto = self.extract_text(documento.ruta_almacenamiento, documento.formato)
            if not texto or not texto.strip():
                raise ValueError("No se pudo extraer texto del documento")
            self.log_info(documento_id, "extraccion_texto", f"Texto extraído: {len(texto)} caracteres")

            # 2. Chunking
            self.log_info(documento_id, "chunking", "Dividiendo texto en fragmentos")
            chunks = self.chunk_text(texto)
            self.log_info(documento_id, "chunking", f"Creados {len(chunks)} fragmentos")

            # 3. Embeddings
            self.log_info(documento_id, "embeddings", "Generando embeddings")
            embeddings = self.generate_embeddings(chunks)
            self.log_info(documento_id, "embeddings", f"Generados {len(embeddings)} embeddings")

            # 4. Almacenar en ChromaDB y PostgreSQL
            self.log_info(documento_id, "almacenamiento", "Guardando fragmentos y vectores")
            self.store_embeddings(documento_id, chunks, embeddings)
            self.log_info(documento_id, "almacenamiento", "Fragmentos guardados correctamente")

            self.update_status(documento_id, "procesado")
            self.log_info(documento_id, "completado", "Procesamiento completado exitosamente")
            return True

        except Exception as e:
            self.update_status(documento_id, "error")
            self.log_error(documento_id, "error_procesamiento", str(e))
            logger.error(f"Error procesando documento {documento_id}: {e}")
            return False

    def process_document_with_ai(self, documento_id: uuid.UUID) -> bool:
        """Procesa documento y luego ejecuta IA (clasificación, resumen, extracción)."""
        from app.services.ai_service import AIService
        from app.services.document_service import DocumentService
        
        # Primero procesamiento básico
        success = self.process_document_sync(documento_id)
        if not success:
            return False
        
        # Luego IA
        try:
            doc_service = DocumentService(self.db)
            ai_service = AIService()
            
            # Obtener texto completo
            from app.models import Fragmento
            fragmentos = self.db.query(Fragmento).filter(Fragmento.documento_id == documento_id).order_by(Fragmento.orden).all()
            texto_completo = "\n\n".join([f.texto for f in fragmentos])
            
            # IA: clasificación, resumen, extracción
            categoria = ai_service.clasificar(texto_completo)
            resumen = ai_service.resumir(texto_completo)
            campos = ai_service.extraer_campos(texto_completo, categoria)
            
            doc_service.update_documento(documento_id, {
                "categoria": categoria,
                "resumen": resumen,
                "campos_extraidos": campos,
                "estado": "procesado"
            })
            self.log_info(documento_id, "ia_completado", f"Categoría: {categoria}")
            return True
            
        except Exception as e:
            self.update_status(documento_id, "error")
            self.log_error(documento_id, "error_ia", str(e))
            logger.error(f"Error en IA para documento {documento_id}: {e}")
            return False

    def update_status(self, documento_id: uuid.UUID, estado: str):
        documento = self.db.query(Documento).filter(Documento.id == documento_id).first()
        if documento:
            documento.estado = estado
            if estado == "procesado":
                from datetime import datetime, timezone
                documento.procesado_en = datetime.now(timezone.utc)
            self.db.commit()

    def log_info(self, documento_id: uuid.UUID, tipo: str, mensaje: str):
        log = LogProcesamiento(documento_id=documento_id, tipo=tipo, mensaje=mensaje)
        self.db.add(log)
        self.db.commit()

    def log_error(self, documento_id: uuid.UUID, tipo: str, mensaje: str):
        log = LogProcesamiento(documento_id=documento_id, tipo=f"error_{tipo}", mensaje=mensaje)
        self.db.add(log)
        self.db.commit()